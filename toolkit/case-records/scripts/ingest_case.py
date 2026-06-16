#!/usr/bin/env python3
"""
Ingest a single case directory into case-records DB.

- Walks case_dir, finds pleading/application files.
- Extracts text per file.
- Parses filename for metadata (doc_type, doc_date, author_role).
- Chunks, embeds, inserts.
"""

import argparse
import hashlib
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import numpy as np

ROOT = Path(os.path.expanduser("~/case-records"))
DB_PATH = ROOT / "db" / "cases_fts.db"
SECRETS = Path(os.path.expanduser("~/.jurisupport/secrets.env"))
try:
    from dotenv import load_dotenv
except ImportError:
    load_dotenv = None

if load_dotenv:
    load_dotenv(SECRETS)

CHUNK_SIZE = 1500
CHUNK_OVERLAP = 300
TRUE_VALUES = {"1", "true", "yes", "y", "on"}
SUPPORTED_EXTS = {".pdf", ".doc", ".docx", ".md", ".txt", ".hwp", ".hwpx"}
TARGET_DOC_CATEGORIES = {"argument", "application"}

ARGUMENT_KEYWORDS = (
    "준비서면",
    "답변서",
    "소장",
    "항소이유서",
    "상고이유서",
    "의견서",
    "변론요지서",
    "참고서면",
    "반박서면",
    "주장서면",
    "청구취지",
    "청구원인",
    "항변",
)
APPLICATION_KEYWORDS = (
    "신청서",
    "항고장",
    "이의신청",
    "보정서",
    "문서제출명령",
    "사실조회",
    "증거신청",
    "석명",
    "감정",
    "검증",
    "증인신청",
    "기록열람",
    "열람등사",
    "소송비용액확정",
    "지급명령신청",
    "가압류",
    "가처분",
)
ADMIN_OR_EVIDENCE_KEYWORDS = (
    "위임계약",
    "소송위임장",
    "영수증",
    "납부",
    "접수증",
    "송달증명",
    "확정증명",
    "기일통지",
    "보정명령",
    "판결문",
    "결정문",
    "갑호증",
    "을호증",
    "증거설명서",
    "계약서",
    "등본",
    "초본",
    "등기",
    "사업자등록",
    "인감",
    "진단서",
    "사진",
    "녹취",
    "카톡",
    "문자",
    "세금계산서",
)

# Filename pattern: {case#}_{idx}_{YYYY.MM.DD}_{doctype}_..._{author}.pdf
FILENAME_RE = re.compile(
    r"^([^_]+)_(\d+)_(\d{4}\.\d{1,2}\.\d{1,2})_([^_]+)(?:_.*?)?(?:_([^_.]+))?\.(pdf|docx?|md|txt|hwp|hwpx)$",
    re.IGNORECASE,
)


def parse_filename(fname: str):
    m = FILENAME_RE.match(fname)
    if not m:
        return None
    return {
        "case_no": m.group(1),
        "seq": m.group(2),
        "doc_date": m.group(3).replace(".", "-"),
        "doc_type": m.group(4),
        "author_role": m.group(5) or "",
    }


def ensure_schema(con: sqlite3.Connection) -> None:
    con.executescript("""
CREATE TABLE IF NOT EXISTS cases (
  case_id TEXT PRIMARY KEY,
  case_name TEXT,
  status TEXT,
  result TEXT,
  court TEXT,
  added_at TEXT DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS documents (
  doc_id TEXT PRIMARY KEY,
  case_id TEXT NOT NULL REFERENCES cases(case_id),
  doc_type TEXT,
  doc_date TEXT,
  author_role TEXT,
  source_file TEXT
);
CREATE TABLE IF NOT EXISTS chunks (
  chunk_id TEXT PRIMARY KEY,
  doc_id TEXT NOT NULL REFERENCES documents(doc_id),
  case_id TEXT NOT NULL REFERENCES cases(case_id),
  chunk_text TEXT NOT NULL,
  embedding BLOB
);
CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(
  chunk_text, chunk_id UNINDEXED, case_id UNINDEXED, doc_id UNINDEXED,
  content='chunks', content_rowid='rowid', tokenize='unicode61'
);
""")
    existing = {r[1] for r in con.execute("PRAGMA table_info(documents)")}
    if "source_kind" not in existing:
        con.execute("ALTER TABLE documents ADD COLUMN source_kind TEXT DEFAULT 'mixed'")
    if "doc_category" not in existing:
        con.execute("ALTER TABLE documents ADD COLUMN doc_category TEXT DEFAULT 'other'")
    con.execute("CREATE INDEX IF NOT EXISTS idx_documents_case_source ON documents(case_id, source_kind)")


def clear_existing_source(con: sqlite3.Connection, case_id: str, source_kind: str) -> None:
    rows = con.execute(
        "SELECT doc_id FROM documents WHERE case_id = ? AND COALESCE(source_kind, 'mixed') = ?",
        (case_id, source_kind),
    ).fetchall()
    if not rows:
        return
    doc_ids = [r[0] for r in rows]
    placeholders = ",".join("?" for _ in doc_ids)
    con.execute(f"DELETE FROM chunks WHERE doc_id IN ({placeholders})", doc_ids)
    con.execute(f"DELETE FROM documents WHERE doc_id IN ({placeholders})", doc_ids)


def text_key(path: Path, doc_type: str = "") -> str:
    return " ".join([str(path.parent), path.stem, doc_type]).lower()


def infer_doc_type(path: Path) -> str:
    key = text_key(path)
    for kw in ARGUMENT_KEYWORDS + APPLICATION_KEYWORDS:
        if kw.lower() in key:
            return kw
    return "기타"


def classify_doc_category(path: Path, doc_type: str) -> str:
    key = text_key(path, doc_type)
    for kw in ARGUMENT_KEYWORDS:
        if kw.lower() in key:
            return "argument"
    for kw in APPLICATION_KEYWORDS:
        if kw.lower() in key:
            return "application"
    for kw in ADMIN_OR_EVIDENCE_KEYWORDS:
        if kw.lower() in key:
            return "other"
    return "other"


def run_text_command(cmd: list[str], timeout: int = 60) -> str:
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, check=False)
    except (OSError, subprocess.SubprocessError):
        return ""
    return p.stdout if p.returncode == 0 else ""


def extract_doc(path: Path) -> str:
    if shutil.which("textutil"):
        text = run_text_command(["textutil", "-convert", "txt", "-stdout", str(path)])
        if text.strip():
            return text
    if shutil.which("antiword"):
        return run_text_command(["antiword", str(path)])
    return ""


def extract_hwp(path: Path) -> str:
    hwpjs = shutil.which("hwpjs")
    if not hwpjs:
        return ""
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "out.md"
        run_text_command([hwpjs, "to-markdown", str(path), "-o", str(out)], timeout=120)
        try:
            return out.read_text(encoding="utf-8", errors="ignore")
        except OSError:
            return ""


def extract_hwpx(path: Path) -> str:
    out = []
    try:
        with zipfile.ZipFile(path) as z:
            names = [n for n in z.namelist() if n.lower().endswith(".xml")]
            for name in names:
                if not (name.startswith("Contents/") or name.startswith("BodyText/") or "section" in name.lower()):
                    continue
                try:
                    root = ET.fromstring(z.read(name))
                except ET.ParseError:
                    continue
                for elem in root.iter():
                    if elem.text and elem.text.strip():
                        out.append(elem.text.strip())
    except (OSError, zipfile.BadZipFile):
        return ""
    return "\n".join(out)


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
        if ext == ".doc":
            return extract_doc(path)
        if ext in (".docx",):
            from docx import Document
            return "\n".join(p.text for p in Document(str(path)).paragraphs)
        if ext in (".md", ".txt"):
            return path.read_text(encoding="utf-8", errors="ignore")
        if ext == ".hwpx":
            return extract_hwpx(path)
        if ext == ".hwp":
            return extract_hwp(path)
        return ""
    except Exception as e:
        print(f"  [warn] extract failed {path.name}: {e}", file=sys.stderr)
        return ""


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP):
    text = text.strip()
    if not text:
        return []
    out = []
    i = 0
    while i < len(text):
        end = min(i + size, len(text))
        out.append(text[i:end])
        if end == len(text):
            break
        i += size - overlap
    return out


def embed_batch(texts: list[str]) -> list[list[float]]:
    from google import genai
    key = os.environ.get("GEMINI_API_KEY")
    if not key:
        raise RuntimeError("GEMINI_API_KEY not set in ~/.jurisupport/secrets.env")
    client = genai.Client(api_key=key)
    out = []
    for i in range(0, len(texts), 100):
        batch = texts[i:i + 100]
        r = client.models.embed_content(model="text-embedding-004", contents=batch)
        out.extend([e.values for e in r.embeddings])
        time.sleep(0.5)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--case-dir", required=True, type=Path)
    ap.add_argument("--case-id", required=True)
    ap.add_argument("--case-name", required=True)
    ap.add_argument("--status", default="진행중")
    ap.add_argument("--result", default="")
    ap.add_argument("--court", default="")
    ap.add_argument(
        "--source-kind",
        choices=("record", "draft", "mixed"),
        default="mixed",
        help="Where this directory came from: record=received case records, draft=office-authored documents.",
    )
    ap.add_argument(
        "--doc-scope",
        choices=("target", "all"),
        default="target",
        help="target indexes only argument/application pleadings; all indexes every supported file.",
    )
    ap.add_argument(
        "--allow-external-embedding",
        action="store_true",
        help="Send case text chunks to Gemini for semantic embeddings.",
    )
    args = ap.parse_args()
    allow_external_embedding = (
        args.allow_external_embedding
        or os.environ.get("CASE_RECORDS_ALLOW_EXTERNAL_EMBEDDING", "").lower() in TRUE_VALUES
    )

    con = sqlite3.connect(DB_PATH)
    ensure_schema(con)
    con.execute(
        "INSERT OR REPLACE INTO cases (case_id, case_name, status, result, court) VALUES (?,?,?,?,?)",
        (args.case_id, args.case_name, args.status, args.result, args.court),
    )
    clear_existing_source(con, args.case_id, args.source_kind)

    # Walk case dir
    docs = []
    for p in args.case_dir.rglob("*"):
        if not p.is_file():
            continue
        if p.suffix.lower() not in SUPPORTED_EXTS:
            continue
        meta = parse_filename(p.name) or {}
        doc_type = meta.get("doc_type") or infer_doc_type(p)
        doc_category = classify_doc_category(p, doc_type)
        if args.doc_scope == "target" and doc_category not in TARGET_DOC_CATEGORIES:
            continue
        text = extract_text(p)
        if not text.strip():
            continue
        rel = str(p.relative_to(args.case_dir))
        digest = hashlib.sha1(f"{args.source_kind}:{rel}".encode("utf-8")).hexdigest()[:16]
        doc_id = f"{args.case_id}__{args.source_kind}__{digest}"
        docs.append({
            "doc_id": doc_id,
            "case_id": args.case_id,
            "doc_type": doc_type,
            "doc_category": doc_category,
            "doc_date": meta.get("doc_date", ""),
            "author_role": meta.get("author_role", ""),
            "source_kind": args.source_kind,
            "source_file": str(p),
            "text": text,
        })

    if not docs:
        print(f"  [warn] no text extracted from {args.case_dir}", file=sys.stderr)
        con.commit(); con.close()
        return

    print(f"  [ingest] {len(docs)} documents found")

    all_chunks = []
    for d in docs:
        con.execute(
            """
            INSERT OR REPLACE INTO documents
              (doc_id, case_id, doc_type, doc_date, author_role, source_file, source_kind, doc_category)
            VALUES (?,?,?,?,?,?,?,?)
            """,
            (
                d["doc_id"], d["case_id"], d["doc_type"], d["doc_date"],
                d["author_role"], d["source_file"], d["source_kind"], d["doc_category"],
            ),
        )
        for j, c in enumerate(chunk_text(d["text"])):
            all_chunks.append({
                "chunk_id": f"{d['doc_id']}__{j:04d}",
                "doc_id": d["doc_id"],
                "case_id": d["case_id"],
                "chunk_text": c,
            })

    if not all_chunks:
        con.commit(); con.close()
        return

    if allow_external_embedding:
        print("  [privacy] External embedding enabled: sending case text chunks to Gemini.")
        print(f"  [ingest] {len(all_chunks)} chunks → embedding")
        embs = embed_batch([c["chunk_text"] for c in all_chunks])
    else:
        print("  [privacy] External embedding disabled: storing local FTS index only.")
        print("  [privacy] Use --allow-external-embedding only after confirming client-data policy.")
        embs = [None] * len(all_chunks)

    for c, e in zip(all_chunks, embs):
        blob = np.array(e, dtype=np.float32).tobytes() if e is not None else None
        con.execute(
            "INSERT OR REPLACE INTO chunks (chunk_id, doc_id, case_id, chunk_text, embedding) VALUES (?,?,?,?,?)",
            (c["chunk_id"], c["doc_id"], c["case_id"], c["chunk_text"], blob),
        )

    con.execute("INSERT INTO chunks_fts(chunks_fts) VALUES('rebuild')")
    con.commit()
    con.close()
    print(f"  [done] case {args.case_id} indexed.")


if __name__ == "__main__":
    main()
